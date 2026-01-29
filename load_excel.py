import shutil
from openpyxl import load_workbook
from datetime import datetime
import tkinter as tk
from tkinter import messagebox

TEMPLATE_FILE = "template/Remote_Access_Template.xlsx"
OUTPUT_FILE = f"output/Remote_Access_Request_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def write_value_next_to_label(ws, label, value, date_format=None):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and str(cell.value).strip() == label:
                target = ws.cell(row=cell.row, column=cell.column + 1)

                for merged in ws.merged_cells.ranges:
                    if target.coordinate in merged:
                        top_left = merged.coord.split(":")[0]
                        ws[top_left] = value
                        if date_format:
                            ws[top_left].number_format = date_format
                        return

                target.value = value
                if date_format:
                    target.number_format = date_format
                return


def generate_excel(doc_ref, version, date_str):
    shutil.copy(TEMPLATE_FILE, OUTPUT_FILE)

    wb = load_workbook(OUTPUT_FILE)
    ws = wb.active

    date_obj = datetime.strptime(date_str, "%d-%b-%Y")

    write_value_next_to_label(ws, "DOC. REF. NO.", doc_ref)
    write_value_next_to_label(ws, "VERSION NO.", version)
    write_value_next_to_label(ws, "DATE:", date_obj, "DD-MMM-YYYY")

    wb.save(OUTPUT_FILE)


# ---------------- GUI ---------------- #

def on_submit():
    doc_ref = doc_ref_entry.get()
    version = version_entry.get()
    date_val = date_entry.get()

    if not doc_ref or not version or not date_val:
        messagebox.showerror("Error", "All fields are required")
        return

    try:
        datetime.strptime(date_val, "%d-%b-%Y")
    except ValueError:
        messagebox.showerror("Error", "Date format must be DD-MMM-YYYY")
        return

    generate_excel(doc_ref, version, date_val)
    messagebox.showinfo("Success", "Excel form generated successfully")
    root.destroy()


root = tk.Tk()
root.title("Generate Access Form")
root.geometry("420x220")
root.resizable(False, False)

tk.Label(root, text="DOC. REF. NO.").pack(pady=5)
doc_ref_entry = tk.Entry(root, width=40)
doc_ref_entry.insert(0, f"IPACK-PAM-{datetime.now().year}-001")
doc_ref_entry.pack()

tk.Label(root, text="VERSION NO.").pack(pady=5)
version_entry = tk.Entry(root, width=40)
version_entry.insert(0, "v1.0")
version_entry.pack()

tk.Label(root, text="DATE (DD-MMM-YYYY)").pack(pady=5)
date_entry = tk.Entry(root, width=40)
date_entry.insert(0, datetime.now().strftime("%d-%b-%Y"))
date_entry.pack()

tk.Button(root, text="Generate", command=on_submit).pack(pady=15)

root.mainloop()
